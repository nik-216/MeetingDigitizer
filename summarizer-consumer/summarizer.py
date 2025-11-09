import re
from datetime import datetime
from typing import List, Dict, Optional
from docx import Document
from transformers import PegasusTokenizer, PegasusForConditionalGeneration
import torch

# ---------------- Timestamp Extraction ----------------
_BASE_DT = None

def extract_timestamp_seconds(line: str) -> Optional[float]:
    global _BASE_DT
    m_dt = re.search(r"Timestamp:\s*([\d-]+\s[\d:]+)", line)
    if m_dt:
        try:
            dt = datetime.strptime(m_dt.group(1), "%Y-%m-%d %H:%M:%S")
            if _BASE_DT is None:
                _BASE_DT = dt
            return (dt - _BASE_DT).total_seconds()
        except Exception:
            pass
    m_secs = re.search(r"Timestamp:\s*([\d.]+)s", line)
    if m_secs:
        try:
            return float(m_secs.group(1))
        except Exception:
            pass
    return None


def format_timecode(seconds: float) -> str:
    seconds = max(0.0, float(seconds))
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    ms = int(round((seconds - int(seconds)) * 1000))
    return f"{h:02d}:{m:02d}:{s:02d}.{ms:03d}"


# ---------------- Readers ----------------
def read_audio_docx(path: str):
    global _BASE_DT
    _BASE_DT = None
    doc = Document(path)
    entries = []
    current_t, speaker, transcript = None, None, None

    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith("Timestamp:"):
            current_t = extract_timestamp_seconds(text)
        elif text.startswith("Speaker:"):
            speaker = text.split("Speaker:")[-1].strip()
        elif text.startswith("Transcript:"):
            transcript = text.split("Transcript:")[-1].strip()
            if current_t is not None and speaker and transcript:
                entries.append({
                    "t": current_t,
                    "type": "audio",
                    "speaker": speaker,
                    "content": transcript
                })
                current_t, speaker, transcript = None, None, None
    return entries


def read_ocr_docx(path: str):
    global _BASE_DT
    _BASE_DT = None
    doc = Document(path)
    entries = []
    current_t, detected = None, None

    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith("Timestamp:"):
            current_t = extract_timestamp_seconds(text)
        elif text.startswith("Detected Text:"):
            detected = text.split("Detected Text:")[-1].strip()
            if current_t is not None and detected:
                entries.append({
                    "t": current_t,
                    "type": "text",
                    "content": detected
                })
                current_t, detected = None, None
    return entries


def read_diagram_docx(path: str):
    global _BASE_DT
    _BASE_DT = None
    doc = Document(path)
    entries = []
    current_t, extracted, inference = None, None, None

    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith("Timestamp:"):
            current_t = extract_timestamp_seconds(text)
        elif text.startswith("Extracted Text:"):
            extracted = text.split("Extracted Text:")[-1].strip()
        elif text.startswith("Inference:"):
            inference = text.split("Inference:")[-1].strip()
            if current_t is not None:
                entries.append({
                    "t": current_t,
                    "type": "diagram",
                    "extracted": extracted,
                    "inference": inference
                })
                current_t, extracted, inference = None, None, None
    return entries


# ---------------- Merge Logic ----------------
def merge_entries(audio_entries, text_entries, diagram_entries, window=5.0):
    all_entries = audio_entries + text_entries + diagram_entries
    all_entries = [e for e in all_entries if isinstance(e.get("t"), (int, float))]
    buckets: Dict[float, List[dict]] = {}
    for e in all_entries:
        key = round(e["t"] / window) * window
        buckets.setdefault(key, []).append(e)
    merged = [{"t": k, "group": buckets[k]} for k in sorted(buckets.keys())]
    return merged


def write_merged_docx(merged: List[dict], output_path: str):
    doc = Document()
    for block in merged:
        doc.add_paragraph(f"{format_timecode(block['t'])}")
        for e in block["group"]:
            if e["type"] == "audio":
                doc.add_paragraph(f"{e['speaker']}: {e['content']}")
            elif e["type"] == "text":
                doc.add_paragraph(f"OCR: {e['content']}")
            elif e["type"] == "diagram":
                doc.add_paragraph(f"Diagram Inference: {e.get('inference','')}")
                if e.get("extracted"):
                    doc.add_paragraph(f"Extracted Text: {e['extracted']}")
        doc.add_paragraph("-" * 80)
    doc.save(output_path)


def merge_docx_files(audio_path: str, ocr_path: str, diagram_path: str, output_path: str, window: float = 5.0):
    audio = read_audio_docx(audio_path)
    text = read_ocr_docx(ocr_path)
    diagram = read_diagram_docx(diagram_path)
    merged = merge_entries(audio, text, diagram, window)
    write_merged_docx(merged, output_path)
    print(f"Merged document saved to: {output_path}")


# ---------------- Cleaning & Chunking ----------------
def clean_text_for_summarization(text: str) -> str:
    text = re.sub(r"\b\d{2}:\d{2}:\d{2}\.\d{3}\b", "", text)
    text = re.sub(r"(OCR:|Diagram Inference:|Extracted Text:|Speaker:|-+)", "", text)
    text = re.sub(r"\[.*?\]", "", text)
    text = re.sub(r"\{.*?\}", "", text)
    text = re.sub(r"[\n\r]+", "\n", text)
    text = re.sub(r"\s{2,}", " ", text)
    return text.strip()


def chunk_text(text, tokenizer, chunk_size=512, overlap=50):
    tokens = tokenizer.encode(text, truncation=False)
    chunks = []
    start = 0
    while start < len(tokens):
        end = min(start + chunk_size, len(tokens))
        chunks.append(tokens[start:end])
        start += chunk_size - overlap
    return chunks


def postprocess_summary(text: str) -> str:
    lines = text.split("\n")
    seen = set()
    cleaned = []
    for line in lines:
        line = line.strip()
        if not line or len(line.split()) < 3:
            continue
        if line.lower() in seen:
            continue
        seen.add(line.lower())
        cleaned.append(line)
    return "\n".join(cleaned)


# ---------------- Summarization ----------------
def summarize_docx(input_path: str, output_path: str, model_dir: str = "pegasus-model", chunk_size: int = 512):
    print("Loading Pegasus model from:", model_dir)
    tokenizer = PegasusTokenizer.from_pretrained(model_dir)
    model = PegasusForConditionalGeneration.from_pretrained(model_dir)
    model.eval()

    doc = Document(input_path)
    raw_text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    cleaned_text = clean_text_for_summarization(raw_text)

    if not cleaned_text.strip():
        raise ValueError("No valid text found in merged DOCX.")

    chunks = chunk_text(cleaned_text, tokenizer, chunk_size)
    print(f"Summarizing {len(chunks)} chunks...")

    summaries = []
    for i, chunk in enumerate(chunks):
        batch = tokenizer.decode(chunk, skip_special_tokens=True)
        inputs = tokenizer([batch], max_length=chunk_size, truncation=True, return_tensors="pt")

        with torch.no_grad():
            summary_ids = model.generate(
                **inputs,
                max_length=128,
                num_beams=4,
                length_penalty=2.0,
                early_stopping=True
            )

        summary_text = tokenizer.decode(summary_ids[0], skip_special_tokens=True)
        summaries.append(summary_text)
        print(f"Chunk {i+1}/{len(chunks)} summarized.")

    final_summary = postprocess_summary("\n\n".join(summaries))

    summary_doc = Document()
    summary_doc.add_heading("Pegasus Summary", level=1)
    summary_doc.add_paragraph(final_summary)
    summary_doc.save(output_path)
    print(f"Summary saved to: {output_path}")


# ---------------- Main ----------------
if __name__ == "__main__":
    merge_docx_files(
        "../output/audio_transcripts.docx",
        "../output/ocr_sentences.docx",
        "../output/diagram_detections.docx",
        "../output/merged_output.docx",
        window=5
    )

    summarize_docx(
        input_path="../output/merged_output.docx",
        output_path="../output/merged_summary.docx",
        model_dir="pegasus-model"
    )
