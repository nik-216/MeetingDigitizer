import json
import time
import base64
import os
from datetime import datetime
from kafka import KafkaConsumer
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn
import threading
from queue import PriorityQueue
import logging

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class MultiTopicKafkaConsumer:
    def __init__(self, bootstrap_servers=['kafka:9092'], output_dir='./output'):
        self.bootstrap_servers = bootstrap_servers
        self.output_dir = output_dir
        
        # Create output directory
        os.makedirs(self.output_dir, exist_ok=True)
        
        # Store messages for each topic separately
        self.audio_messages = []
        self.ocr_messages = []
        self.diagram_messages = []
        
        # Fixed document filenames for each topic
        self.audio_doc_filename = os.path.join(self.output_dir, 'audio_transcripts.docx')
        self.ocr_doc_filename = os.path.join(self.output_dir, 'ocr_sentences.docx')
        self.diagram_doc_filename = os.path.join(self.output_dir, 'diagram_detections.docx')
        
        # Topics to consume from
        self.topics = ['audio-transcripts', 'ocr-sentences', 'diagram-detections']
        
        # Consumer configuration
        self.consumer = KafkaConsumer(
            *self.topics,
            bootstrap_servers=self.bootstrap_servers,
            auto_offset_reset='latest',  # Change to 'earliest' to consume from beginning
            enable_auto_commit=True,
            group_id='multi-topic-consumer-group',
            value_deserializer=lambda x: json.loads(x.decode('utf-8'))
        )
        
        # Track processed messages count
        self.processed_count = 0
        
    def get_timestamp_from_message(self, message_data, topic):
        """Extract timestamp from message based on topic"""
        try:
            if topic == 'audio-transcripts':
                return message_data.get('kafka_timestamp', time.time())
            elif topic == 'ocr-sentences':
                return message_data.get('detection_timestamp', time.time())
            elif topic == 'diagram-detections':
                return message_data.get('detection_timestamp', time.time())
            else:
                return time.time()
        except:
            return time.time()
    
    def create_audio_document(self):
        """Create Word document for audio transcripts"""
        doc = Document()
        doc.add_heading('Audio Transcripts', 0)
        
        # Add document info
        p = doc.add_paragraph()
        p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").bold = True
        p = doc.add_paragraph(f"Total Messages: {len(self.audio_messages)}")
        doc.add_page_break()
        
        # Sort messages by timestamp
        sorted_messages = sorted(self.audio_messages, key=lambda x: x[0])
        
        # Process each message
        for timestamp, message_data in sorted_messages:
            self.format_audio_transcript_to_doc(doc, message_data)
        
        return doc
    
    def create_ocr_document(self):
        """Create Word document for OCR sentences"""
        doc = Document()
        doc.add_heading('OCR Text Detections', 0)
        
        # Add document info
        p = doc.add_paragraph()
        p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").bold = True
        p = doc.add_paragraph(f"Total Messages: {len(self.ocr_messages)}")
        doc.add_page_break()
        
        # Sort messages by timestamp
        sorted_messages = sorted(self.ocr_messages, key=lambda x: x[0])
        
        # Process each message
        for timestamp, message_data in sorted_messages:
            self.format_ocr_sentence_to_doc(doc, message_data)
        
        return doc
    
    def create_diagram_document(self):
        """Create Word document for diagram detections"""
        doc = Document()
        doc.add_heading('Diagram Detections', 0)
        
        # Add document info
        p = doc.add_paragraph()
        p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").bold = True
        p = doc.add_paragraph(f"Total Messages: {len(self.diagram_messages)}")
        doc.add_page_break()
        
        # Sort messages by timestamp
        sorted_messages = sorted(self.diagram_messages, key=lambda x: x[0])
        
        # Process each message
        for timestamp, message_data in sorted_messages:
            self.format_diagram_detection_to_doc(doc, message_data)
        
        return doc
    
    def add_separator_to_doc(self, doc):
        """Add a separator line to the document"""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("─" * 80)
        run.font.color.rgb = None  # Default color
    
    def format_audio_transcript_to_doc(self, doc, data):
        """Format audio transcript message for Word document"""
        doc.add_heading(f'🎤 Audio Transcript', level=2)
        
        # Add timestamp
        timestamp = data.get('kafka_timestamp', time.time())
        readable_time = datetime.fromtimestamp(timestamp).strftime('%Y-%m-%d %H:%M:%S')
        p = doc.add_paragraph()
        p.add_run(f"Timestamp: {readable_time}").bold = True
        
        # Add time range
        start_time = data.get('start', 'N/A')
        end_time = data.get('end', 'N/A')
        p = doc.add_paragraph(f"Time Range: {start_time}s - {end_time}s")
        
        # Add speaker
        speaker = data.get('speaker', 'Unknown')
        p = doc.add_paragraph(f"Speaker: {speaker}")
        
        # Add transcript text
        text = data.get('text', '')
        p = doc.add_paragraph()
        p.add_run("Transcript: ").bold = True
        p.add_run(text)
        
        self.add_separator_to_doc(doc)
    
    def format_ocr_sentence_to_doc(self, doc, data):
        """Format OCR sentence message for Word document"""
        doc.add_heading(f'📄 OCR Text Detection', level=2)
        
        # Add timestamp
        timestamp = data.get('detection_timestamp', time.time())
        readable_time = datetime.fromtimestamp(timestamp).strftime('%Y-%m-%d %H:%M:%S')
        p = doc.add_paragraph()
        p.add_run(f"Timestamp: {readable_time}").bold = True
        
        # Add frame info
        frame_id = data.get('frame_id', 'N/A')
        p = doc.add_paragraph(f"Frame ID: {frame_id}")
        
        # Add line number
        line_number = data.get('line_number', 'N/A')
        p = doc.add_paragraph(f"Line Number: {line_number}")
        
        # Add confidence and word count
        confidence = data.get('confidence', 0)
        word_count = data.get('word_count', 0)
        p = doc.add_paragraph(f"Confidence: {confidence:.2f} | Word Count: {word_count}")
        
        # Add detected text
        sentence = data.get('sentence', '')
        p = doc.add_paragraph()
        p.add_run("Detected Text: ").bold = True
        p.add_run(sentence)
        
        self.add_separator_to_doc(doc)
    
    def format_diagram_detection_to_doc(self, doc, data):
        """Format diagram detection message for Word document"""
        doc.add_heading(f'Diagram Detection', level=2)
        
        # Add timestamp
        timestamp = data.get('detection_timestamp', time.time())
        readable_time = datetime.fromtimestamp(timestamp).strftime('%Y-%m-%d %H:%M:%S')
        p = doc.add_paragraph()
        p.add_run(f"Timestamp: {readable_time}").bold = True
        
        # Add frame info
        frame_id = data.get('frame_id', 'N/A')
        diagram_type = data.get('diagram_type', 'Unknown')
        p = doc.add_paragraph(f"Frame ID: {frame_id} | Type: {diagram_type}")
        
        # Add bounding box and dimensions
        bbox = data.get('bbox', [])
        area = data.get('area', 0)
        aspect_ratio = data.get('aspect_ratio', 0)
        p = doc.add_paragraph(f"Bounding Box: {bbox}")
        p = doc.add_paragraph(f"Area: {area} | Aspect Ratio: {aspect_ratio:.2f}")
        
        # Add confidence
        confidence = data.get('confidence', 0)
        p = doc.add_paragraph(f"Confidence: {confidence:.2f}")
        
        # Add extracted text
        extracted_text = data.get('extracted_text', [])
        if extracted_text:
            p = doc.add_paragraph()
            p.add_run("Extracted Text: ").bold = True
            p.add_run(str(extracted_text))
        
        # Embed diagram image directly in document
        diagram_b64 = data.get('diagram_image')
        if diagram_b64:
            try:
                # Decode base64 image
                image_data = base64.b64decode(diagram_b64)
                
                # Create temporary image file
                temp_filename = f"temp_diagram_{frame_id}_{int(timestamp)}.png"
                temp_filepath = os.path.join(self.output_dir, temp_filename)
                
                # Save temporary image
                with open(temp_filepath, 'wb') as f:
                    f.write(image_data)
                
                # Add image to document
                p = doc.add_paragraph()
                p.add_run("Diagram Image:").bold = True
                doc.add_picture(temp_filepath, width=Inches(5))
                
                # Clean up temporary file
                os.remove(temp_filepath)
                
                logger.info(f"Embedded diagram image for frame {frame_id}")
                
            except Exception as e:
                logger.error(f"Error embedding diagram image: {e}")
                p = doc.add_paragraph(f"Diagram Image: Error loading image - {str(e)}")
        
        self.add_separator_to_doc(doc)
    
    def save_all_documents(self):
        """Save all topic documents (overwrite existing)"""
        saved_files = []
        
        try:
            # Save audio transcripts document
            if self.audio_messages:
                audio_doc = self.create_audio_document()
                audio_doc.save(self.audio_doc_filename)
                saved_files.append(self.audio_doc_filename)
                logger.info(f"Audio transcripts saved: {self.audio_doc_filename}")
            
            # Save OCR sentences document
            if self.ocr_messages:
                ocr_doc = self.create_ocr_document()
                ocr_doc.save(self.ocr_doc_filename)
                saved_files.append(self.ocr_doc_filename)
                logger.info(f"OCR sentences saved: {self.ocr_doc_filename}")
            
            # Save diagram detections document
            if self.diagram_messages:
                diagram_doc = self.create_diagram_document()
                diagram_doc.save(self.diagram_doc_filename)
                saved_files.append(self.diagram_doc_filename)
                logger.info(f"Diagram detections saved: {self.diagram_doc_filename}")
            
            return saved_files
            
        except Exception as e:
            logger.error(f"Error saving documents: {e}")
            return saved_files
    
    def consume_messages(self, save_interval=30):
        """
        Main consumer loop - saves each topic to separate Word files
        
        Args:
            save_interval: Interval in seconds to rebuild and save documents
        """
        logger.info(f"Starting consumer for topics: {self.topics}")
        logger.info(f"Output directory: {self.output_dir}")
        logger.info("Each topic will be saved to separate Word files:")
        logger.info(f"  - Audio: {self.audio_doc_filename}")
        logger.info(f"  - OCR: {self.ocr_doc_filename}")
        logger.info(f"  - Diagrams: {self.diagram_doc_filename}")
        
        last_save_time = time.time()
        
        try:
            for message in self.consumer:
                try:
                    topic = message.topic
                    message_data = message.value
                    
                    # Get timestamp for ordering
                    timestamp = self.get_timestamp_from_message(message_data, topic)
                    
                    # Add message to appropriate collection
                    if topic == 'audio-transcripts':
                        self.audio_messages.append((timestamp, message_data))
                    elif topic == 'ocr-sentences':
                        self.ocr_messages.append((timestamp, message_data))
                    elif topic == 'diagram-detections':
                        self.diagram_messages.append((timestamp, message_data))
                    
                    self.processed_count += 1
                    logger.info(f"Received message #{self.processed_count} from {topic}")
                    
                    # Save documents periodically (overwrite)
                    current_time = time.time()
                    if current_time - last_save_time >= save_interval:
                        self.save_all_documents()
                        last_save_time = current_time
                
                except Exception as e:
                    logger.error(f"Error processing message: {e}")
                    continue
        
        except KeyboardInterrupt:
            logger.info("Consumer interrupted by user")
        
        finally:
            # Final save with all collected messages
            saved_files = self.save_all_documents()
            logger.info(f"Final documents saved: {saved_files}")
            logger.info(f"Total messages processed: {self.processed_count}")
            logger.info(f"Audio messages: {len(self.audio_messages)}")
            logger.info(f"OCR messages: {len(self.ocr_messages)}")
            logger.info(f"Diagram messages: {len(self.diagram_messages)}")
            
            # Close consumer
            self.consumer.close()
            logger.info("Consumer closed")

def main():
    """Main function to run the consumer"""
    # Configuration
    KAFKA_SERVERS = ['kafka:9092']  # Update with your Kafka servers
    OUTPUT_DIR = './output'         # Update with your desired output directory
    
    # Create consumer instance
    consumer = MultiTopicKafkaConsumer(
        bootstrap_servers=KAFKA_SERVERS,
        output_dir=OUTPUT_DIR
    )
    
    # Start consuming
    logger.info("Starting Kafka consumer...")
    consumer.consume_messages(
        save_interval=30    # Rebuild and save documents every 30 seconds
    )

if __name__ == "__main__":
    main()